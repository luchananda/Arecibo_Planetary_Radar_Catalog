#!/usr/bin/env python3
"""Merge core CW/DD/Date rows with cross-sheet + cell-comment notes into the
final output workbook requested: one row per (object, date) attempt, with
Object, Year, Date, CW, CW files, Delay-Doppler, DD setup+files, Notes.
"""
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

PROJ = '/Users/floridaspaceinstitute/Documents/Radar/Catalogue4Webpage'

core = json.load(open(f'{PROJ}/scratch_session2/braincell_core_rows.json'))
notes_by_object = json.load(open(f'{PROJ}/scratch_session2/braincell_notes_by_object.json'))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Detections'

headers = ['Object', 'Year', 'Date', 'CW', 'CW files', 'Delay-Doppler', 'DD setup & files', 'Notes']
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='2F5496')
    cell.alignment = Alignment(wrap_text=True, vertical='center')
ws.freeze_panes = 'A2'

for r in core:
    notes = notes_by_object.get(r['object'], '')
    ws.append([
        r['object'], r['year'], r['date'],
        r['cw'], r['cw_files'],
        r['dd'], r['dd_detail'],
        notes,
    ])

widths = [16, 8, 10, 6, 30, 14, 30, 90]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=(cell.column_letter in ('E', 'G', 'H')))

# --- summary sheet ---
ws2 = wb.create_sheet('Summary')
ws2.append(['Metric', 'Value'])
ws2.cell(row=1, column=1).font = Font(bold=True)
ws2.cell(row=1, column=2).font = Font(bold=True)
objs = sorted(set(r['object'] for r in core))
rows_with_notes = sum(1 for r in core if notes_by_object.get(r['object']))
summary = [
    ('Total object-date rows', len(core)),
    ('Distinct objects', len(objs)),
    ('Objects with at least one note', len(notes_by_object)),
    ('Rows carrying a note', rows_with_notes),
    ('Source sheets parsed for CW/DD rows', '1950-1999, 2000-2005, 2006-2010, 2011-2015, 2016-2020, Comets, "Named" objects'),
    ('Other sheets mined for notes', 'Not_processed, No_Detection, No Data FoundNo Folder, Re-check, Compare, Objects with data, Successful_Detection, ALL OBJECTS THAT HAVE A FOLDER, MISSING, Sheet17'),
    ('Cell comments extracted (real text)', 669),
    ('Cell comments attributed to an object', 653),
]
for row in summary:
    ws2.append(row)
ws2.column_dimensions['A'].width = 40
ws2.column_dimensions['B'].width = 70

out_path = f'{PROJ}/Data_Project_brain.cell.killer_extracted.xlsx'
wb.save(out_path)
print(f'Wrote {out_path}: {len(core)} rows, {len(objs)} objects, {len(notes_by_object)} objects with notes')
