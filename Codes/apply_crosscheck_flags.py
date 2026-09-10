#!/usr/bin/env python3
"""Turn the D-findings (117 objects marked detected on the dashboard, but the
brain.cell.killer Compare/No_Detection tracking sheets say 'no detection')
plus the 2010 JL33 Goldstone case into:
  1. a 'Re-check (detection flag)' sheet in the extracted workbook, with a
     blank Resolution column for manual triage
  2. an appended note on the Detections sheet's Notes column for those objects
  3. an appended note in FINAL_TABLE_all.json's 'notes' field (shows on the
     live dashboard card hover) for those objects
"""
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

PROJ = '/Users/floridaspaceinstitute/Documents/Radar/Catalogue4Webpage'

findings = json.load(open(f'{PROJ}/scratch_session2/crosscheck_findings.json'))
d_findings = findings['D']

JL33_NOTE = ('SBDB / brain.cell.killer notes indicate this observation was made at Goldstone, '
             'not Arecibo - no AO data found for this object. Flagged for review: should this '
             'object remain marked as an Arecibo detection?')

flags = [(x['target'], 'Marked detected, but brain.cell.killer\'s own Compare/No_Detection '
          'tracking sheets say "no detection"', x['notes_excerpt']) for x in d_findings]
flags.append(('2010 JL33', 'Marked detected, but no CW/DD file is ever marked available for it '
              'in brain.cell.killer', JL33_NOTE))

# ---------- 1 & 2: extracted workbook ----------
wb = openpyxl.load_workbook(f'{PROJ}/Data_Project_brain.cell.killer_extracted.xlsx')
if 'Dashboard cross-check' in wb.sheetnames:
    del wb['Dashboard cross-check']
if 'Re-check (detection flag)' in wb.sheetnames:
    del wb['Re-check (detection flag)']
ws = wb.create_sheet('Re-check (detection flag)')
headers = ['Object', 'Flag reason', 'Detail', 'Resolution (fill in)']
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='B22222')
    cell.alignment = Alignment(wrap_text=True, vertical='center')
ws.freeze_panes = 'A2'
for target, reason, detail in flags:
    ws.append([target, reason, detail, ''])
widths = [14, 45, 90, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)

det_ws = wb['Detections']
obj_col_idx = {}
notes_col_idx = None
for c in range(1, det_ws.max_column + 1):
    h = det_ws.cell(row=1, column=c).value
    if h == 'Object':
        obj_col = c
    if h == 'Notes':
        notes_col_idx = c

flag_by_target_norm = {}
import re
def norm(s):
    return re.sub(r'[\s\-]', '', str(s or '')).lower()
for target, reason, detail in flags:
    flag_by_target_norm[norm(target)] = f'[dashboard cross-check] {reason}: {detail}'

updated_rows = 0
for r in range(2, det_ws.max_row + 1):
    obj = det_ws.cell(row=r, column=obj_col).value
    flag_text = flag_by_target_norm.get(norm(obj))
    if not flag_text:
        continue
    cell = det_ws.cell(row=r, column=notes_col_idx)
    existing = (cell.value or '').strip()
    if flag_text in existing:
        continue
    cell.value = (existing + ' | ' + flag_text) if existing else flag_text
    updated_rows += 1

wb.save(f'{PROJ}/Data_Project_brain.cell.killer_extracted.xlsx')
print(f'Wrote Re-check sheet ({len(flags)} rows), updated Notes on {updated_rows} Detections rows')

# ---------- 3: FINAL_TABLE_all.json ----------
dash = json.load(open(f'{PROJ}/FINAL_TABLE_all.json'))
json_updated = 0
for r in dash:
    for key in (r.get('target'), r.get('name'), r.get('designation')):
        if key and norm(key) in flag_by_target_norm:
            tag = f'[cross-check flag] {flag_by_target_norm[norm(key)]}'
            existing = (r.get('notes') or '').strip()
            if tag not in existing:
                r['notes'] = (existing + ' | ' + tag) if existing else tag
                json_updated += 1
            break

json.dump(dash, open(f'{PROJ}/FINAL_TABLE_all.json', 'w'), indent=1, ensure_ascii=False)
print(f'Updated notes for {json_updated} objects in FINAL_TABLE_all.json')
